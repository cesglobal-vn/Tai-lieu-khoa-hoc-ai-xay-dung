# -*- coding: utf-8 -*-
"""Boc khoi luong phan tho nha pho tu 03_mo-ta-tu-ban-ve.json -> 15_boq-phan-tho-nha-7x20m.xlsx
Chay lai: python boc-khoi-luong.py (cung thu muc voi file JSON va file 14_).
Moi so tinh tren so day du (mm), chi lam tron khi hien thi (number_format 0.00).
"""
import json, os
from collections import Counter, defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.comments import Comment

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "03_mo-ta-tu-ban-ve.json")
SCHED = os.path.join(HERE, "14_bang-thong-ke-cua-nha-7x20m.xlsx")
OUT = os.path.join(HERE, "15_boq-phan-tho-nha-7x20m.xlsx")
THIEU = "[THIẾU KÍCH THƯỚC - CẦN KỸ SƯ XÁC MINH]"
CHO = "[Chờ bổ sung]"

d = json.load(open(SRC, encoding="utf-8"))
cd = d["cao_do"]
mb = d["mat_bang"]
cua_tab = d["cua"]
t_san = d["thong_so"]["Sàn BTCT dày"]  # mm

# Tang tren mat bang -> (cao do san, cao do san tang tren, mat bang tang tren)
# Gia thiet: mat bang "MAI" nam o cao do "ST", tang tren la "TUM" (mai bang M3).
TANG = [("T1", "T1", "T2", "T2"), ("T2", "T2", "T3", "T3"),
        ("T3", "T3", "ST", "MAI"), ("MAI", "ST", "TUM", None)]
GT = []  # gia thiet

def dims(r):
    x1, y1, x2, y2 = r
    dx, dy = abs(x2 - x1), abs(y2 - y1)
    return max(dx, dy), min(dx, dy)

def inside(r, rects):
    x1, y1, x2, y2 = r
    return any(x1 >= a and y1 >= b and x2 <= c and y2 <= e for a, b, c, e in rects)

def slab_above(key_above):
    if key_above is None:  # tang MAI: mai bang tum M3
        return [m["hcn"] for m in mb["MAI"]["mai"] if m["ma"] == "M3"]
    return mb[key_above]["san"]

rows = []  # dict: nhom, ma, ten, dv, dg (dien giai), kl (so day du), nguon, note

# ---------- (1) Tuong xay 200 ----------
for k, lv, lv_up, k_up in TANG:
    H = cd[lv_up] - cd[lv]
    above = slab_above(k_up)
    segs = defaultdict(list)  # chieu cao -> chieu dai doan
    for i, w in enumerate(mb[k]["tuong"]):
        L, t = dims(w)
        if t != 200:
            continue
        h = H - t_san if inside(w, above) else H
        segs[h].append(L)
    ops = [c for c in mb[k]["cua"] if dims(c["hcn"])[1] == 200]
    Lop = sum(dims(c["hcn"])[0] for c in ops)
    start = len(rows)
    for h, Ls in sorted(segs.items()):
        Ltot = sum(Ls)
        rows.append(dict(nhom=1, ma=f"TX-{k}", ten=f"Tường xây 200 tầng {k}: phần tường đặc, cao {h} mm",
                         dv="m2", dg=f"({'+'.join(map(str, Ls))}) mm x {h} mm",
                         kl=Ltot * h / 1e6, nguon=f"mat_bang.{k}.tuong (dày 200); cao_do.{lv}->{lv_up}"))
    # phan tuong tai vi tri lo cua: chieu dai lo x chieu cao tuong, roi tru dien tich cua
    hmap = defaultdict(list)
    for c in ops:
        h = H - t_san if inside(c["hcn"], above) else H
        hmap[h].append(dims(c["hcn"])[0])
    for h, Ls in sorted(hmap.items()):
        rows.append(dict(nhom=1, ma=f"TX-{k}", ten=f"Tường xây 200 tầng {k}: cộng phần tường tại vị trí lỗ cửa",
                         dv="m2", dg=f"({'+'.join(map(str, Ls))}) mm x {h} mm",
                         kl=sum(Ls) * h / 1e6, nguon=f"mat_bang.{k}.cua (lỗ trong tường 200)"))
    cnt = Counter(c["ma"] for c in ops)
    for ma, n in sorted(cnt.items()):
        r_, c_ = cua_tab[ma]["rong"], cua_tab[ma]["cao"]
        rows.append(dict(nhom=1, ma=f"TX-{k}", ten=f"Tường xây 200 tầng {k}: trừ lỗ cửa {ma}",
                         dv="m2", dg=f"-{r_} x {c_} mm x {n}", kl=-r_ * c_ * n / 1e6,
                         nguon=f"mat_bang.{k}.cua ({ma}); cua.{ma}.rong/cao"))
    rows.append(dict(nhom=1, ma=f"TX-{k}", ten=f"Cộng tường xây 200 tầng {k} (diện tích)", dv="m2",
                     dg="Tổng các dòng trên", kl=sum(r["kl"] for r in rows[start:]), nguon="", sub=(start, len(rows))))
    rows.append(dict(nhom=1, ma=f"TX-{k}", ten=f"Tường xây 200 tầng {k} (thể tích)", dv="m3",
                     dg="Diện tích x 0.2 m", kl=rows[-1]["kl"] * 0.2, nguon="", vol_of=len(rows) - 1))
GT.append(f"Tường xây 200: chỉ lấy hình chữ nhật trong mat_bang.*.tuong có cạnh ngắn = 200 mm. Tường 100 (vách WC, vách phòng) không tính. Lan can (mat_bang.*.lan_can) không tính vì file không ghi vật liệu.")
GT.append(f"Chiều cao tường = chênh cao độ tầng (mục cao_do) trừ sàn {t_san} mm nếu phía trên có sàn BTCT; chỗ không có sàn phía trên (dưới mái tôn M1, M2) lấy đủ chênh cao độ. Không có dữ liệu dầm nên không trừ dầm: {THIEU}.")
GT.append("Mặt bằng 'MAI' gán cho cao độ ST (10400) tới TUM (13400). Phần tường hồi tầng 3 nằm trên cao độ ST dưới mái tôn dốc M1 chưa tính: file không ghi hướng dốc. " + THIEU)
GT.append("Đoạn tường trên mặt bằng đã cắt tại lỗ cửa, nên diện tích = phần tường đặc + phần tại lỗ cửa (dài lỗ x cao tường) - diện tích cửa (rộng x cao theo mục cua). Tường dừng tại mép cột, cột tính riêng.")

# ---------- (2) Be tong san ----------
SLABS = [("T1", "Sàn tầng 1 (cốt ±0.000)", mb["T1"]["san"], mb["T1"]["lo_san"], "mat_bang.T1.san"),
         ("T2", "Sàn tầng 2", mb["T2"]["san"], mb["T2"]["lo_san"], "mat_bang.T2.san/lo_san"),
         ("T3", "Sàn tầng 3", mb["T3"]["san"], mb["T3"]["lo_san"], "mat_bang.T3.san/lo_san"),
         ("ST", "Sàn sân thượng", mb["MAI"]["san"], mb["MAI"]["lo_san"], "mat_bang.MAI.san/lo_san"),
         ("TUM", "Sàn mái tum (M3 BTCT)", [m["hcn"] for m in mb["MAI"]["mai"] if m["ma"] == "M3"], [], "mat_bang.MAI.mai.M3")]
for k, ten, san, lo, src in SLABS:
    for r in san:
        a, b = abs(r[2] - r[0]), abs(r[3] - r[1])
        rows.append(dict(nhom=2, ma=f"BTS-{k}", ten=f"Bê tông {ten}", dv="m3",
                         dg=f"{a} x {b} x {t_san} mm", kl=a * b * t_san / 1e9, nguon=f"{src}; thong_so.Sàn BTCT dày"))
    for r in lo:
        a, b = abs(r[2] - r[0]), abs(r[3] - r[1])
        rows.append(dict(nhom=2, ma=f"BTS-{k}", ten=f"Bê tông {ten}: trừ lỗ sàn (thang)", dv="m3",
                         dg=f"-{a} x {b} x {t_san} mm", kl=-a * b * t_san / 1e9, nguon=src))
GT.append(f"Mọi sàn lấy dày {t_san} mm (thong_so 'Sàn BTCT dày'), kể cả sàn tầng 1 và mái tum M3. Sàn tầng 1 có thể là sàn nền trên đất: Cần xác nhận lại.")
GT.append("Diện tích sàn lấy theo hình chữ nhật mục san (phủ cả trên tường, cột), chỉ trừ lo_san. Không trừ phần cột xuyên sàn vì cột đã tính tới đáy sàn.")

# ---------- (3) Be tong cot ----------
for k, lv, lv_up, k_up in TANG:
    H = cd[lv_up] - cd[lv]
    above = slab_above(k_up)
    grp = defaultdict(int)
    for c in mb[k]["cot"]:
        a, b = abs(c[2] - c[0]), abs(c[3] - c[1])
        h = H - t_san if inside(c, above) else H
        grp[(a, b, h)] += 1
    for (a, b, h), n in sorted(grp.items()):
        rows.append(dict(nhom=3, ma=f"BTC-{k}", ten=f"Bê tông cột {a}x{b} tầng {k}", dv="m3",
                         dg=f"{a} x {b} x {h} mm x {n} cột", kl=a * b * h * n / 1e9,
                         nguon=f"mat_bang.{k}.cot; cao_do.{lv}->{lv_up}"))
GT.append("Cột tính từ mặt sàn tầng tới đáy sàn tầng trên (chênh cao độ - 120). Cột không có sàn phía trên (cột trục đầu nhà tầng 2 đỡ mái tôn M2, cột tầng 3 dưới mái M1) lấy đủ chênh cao độ. Cột dưới cốt ±0.000 (cổ cột, móng) không có trong file: " + THIEU)
GT.append("Cột pergola (pergola_cot) là khung thép theo mục vat_lieu, không tính bê tông.")

# ---------- (4) Cua ----------
plan_cnt = Counter()
where = defaultdict(list)
for k in mb:
    for c in mb[k]["cua"]:
        plan_cnt[c["ma"]] += 1
        where[c["ma"]].append(k)
        w = dims(c["hcn"])[0]
        if w != cua_tab[c["ma"]]["rong"]:
            GT.append(f"LỆCH: cửa {c['ma']} tầng {k} rộng trên mặt bằng {w} mm khác bảng cửa {cua_tab[c['ma']]['rong']} mm")
for ma in sorted(cua_tab, key=lambda m: (m[0] != "D", m)):
    t = cua_tab[ma]
    n = plan_cnt[ma]
    dist = ", ".join(f"{k}:{v}" for k, v in Counter(where[ma]).items())
    rows.append(dict(nhom=4, ma=ma, ten=f"{t['ten']} {t['rong']}x{t['cao']}", dv="bộ",
                     dg=f"Đếm trên mặt bằng ({dist})", kl=n, nguon=f"mat_bang.*.cua ma={ma}"))
    rows.append(dict(nhom=4, ma=ma, ten=f"{t['ten']}: diện tích", dv="m2",
                     dg=f"{t['rong']} x {t['cao']} mm x {n}", kl=t["rong"] * t["cao"] * n / 1e6,
                     nguon=f"cua.{ma}.rong/cao"))
GT.append("Số bộ cửa đếm trên mặt bằng (mat_bang.*.cua); kích thước lấy từ mục cua. Đối chiếu ở sheet 'Kiem cheo cua'.")
GT.append("Tên file ghi 'nhà 7x20m' là kích thước lô đất (ranh_dat 7000 x 20000); khối nhà trong file là 5000 x 14000 mm.")
GT.append("Hao hụt (%) và đơn giá để trống (" + CHO + "); công thức Excel tự tính khi điền.")

# ---------- Kiem cheo voi file 14 ----------
wb14 = openpyxl.load_workbook(SCHED, data_only=False)
ws14 = wb14.active
x14 = {}
for r in ws14.iter_rows(min_row=6, values_only=True):
    if r[0] and r[0] in cua_tab:
        x14[r[0]] = dict(rong=r[2], cao=r[3], n=r[5])

# ---------- Xuat Excel ----------
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "BOQ"
B = Font(bold=True)
YEL = PatternFill("solid", fgColor="FFF2CC")
ws.append(["BẢNG KHỐI LƯỢNG PHẦN THÔ NHÀ PHỐ LÔ 7x20M, CÔNG TY CES AI"])
ws.append([f"Nguồn: {os.path.basename(SRC)} (đơn vị mm). Script: boc-khoi-luong.py. Khối lượng hiển thị 2 số lẻ, tính trên số đầy đủ."])
ws.append([])
H = ["STT", "Mã", "Tên công việc", "Đơn vị", "Diễn giải", "KL thiết kế", "Hao hụt (%)",
     "KL hao hụt", "KL mua", "Đơn giá", "Thành tiền", "Nguồn", "Ghi chú"]
ws.append(H)
for c in ws[4]:
    c.font = B
NHOM = {1: "TƯỜNG XÂY 200", 2: "BÊ TÔNG SÀN", 3: "BÊ TÔNG CỘT", 4: "CỬA ĐI, CỬA SỔ"}
excel_row = {}
tot_cells = {}
stt = 0
cur = None
grp_rows = defaultdict(list)
for i, r in enumerate(rows):
    if r["nhom"] != cur:
        cur = r["nhom"]
        ws.append([f"({cur})", "", NHOM[cur]])
        ws.cell(ws.max_row, 3).font = B
    stt += 1
    rr = ws.max_row + 1
    excel_row[i] = rr
    if "sub" in r:
        a, b = r["sub"]
        kl = f"=SUM(F{excel_row[a]}:F{excel_row[b - 1]})"
    elif "vol_of" in r:
        kl = f"=F{excel_row[r['vol_of']]}*0.2"
    else:
        kl = r["kl"]
    ws.append([stt, r["ma"], r["ten"], r["dv"], r["dg"], kl, None,
               f"=F{rr}*G{rr}", f"=F{rr}*(1+G{rr})", None, f"=I{rr}*J{rr}", r["nguon"],
               f"Hao hụt, đơn giá: {CHO}"])
    ws.cell(rr, 7).fill = YEL
    ws.cell(rr, 10).fill = YEL
    ws.cell(rr, 7).comment = Comment(CHO, "boc-khoi-luong")
    for col in (6, 8, 9):
        ws.cell(rr, col).number_format = "#,##0.00"
    ws.cell(rr, 11).number_format = "#,##0"
    if "sub" in r or "vol_of" in r:
        for col in range(1, 14):
            ws.cell(rr, col).font = B
    grp_rows[(r["nhom"], r["dv"], "sub" in r or "vol_of" in r or r["nhom"] != 1)].append(rr)

# Tong nhom
ws.append([])
ws.append(["", "", "TỔNG HỢP"])
ws.cell(ws.max_row, 3).font = B
summ = [(1, "m2", "Tổng tường xây 200 (diện tích)"), (1, "m3", "Tổng tường xây 200 (thể tích)"),
        (2, "m3", "Tổng bê tông sàn"), (3, "m3", "Tổng bê tông cột"),
        (4, "bộ", "Tổng số bộ cửa"), (4, "m2", "Tổng diện tích cửa")]
for g, dv, ten in summ:
    rs = grp_rows[(g, dv, True)]
    f = "=" + "+".join(f"F{x}" for x in rs)
    ws.append(["", "", ten, dv, "", f])
    ws.cell(ws.max_row, 6).number_format = "#,##0.00"
    ws.cell(ws.max_row, 3).font = B
for col, w in zip("ABCDEFGHIJKLM", [5, 9, 50, 7, 48, 13, 11, 11, 11, 11, 13, 42, 26]):
    ws.column_dimensions[col].width = w

# Sheet kiem cheo cua
wk = wb.create_sheet("Kiem cheo cua")
wk.append(["Mã", "Đếm trên mặt bằng", "Mục cua (JSON) so_luong", "File 14_ Số lượng",
           "Rộng JSON", "Rộng file 14", "Cao JSON", "Cao file 14", "Kết quả"])
for c in wk[1]:
    c.font = B
lech = []
for ma in sorted(cua_tab, key=lambda m: (m[0] != "D", m)):
    t = cua_tab[ma]
    x = x14.get(ma, {})
    ok = (plan_cnt[ma] == t["so_luong"] == x.get("n") and t["rong"] == x.get("rong") and t["cao"] == x.get("cao"))
    if not ok:
        lech.append(ma)
    wk.append([ma, plan_cnt[ma], t["so_luong"], x.get("n"), t["rong"], x.get("rong"), t["cao"], x.get("cao"),
               "Khớp" if ok else "LỆCH: cần kỹ sư xác minh"])
extra = set(plan_cnt) - set(cua_tab)
for ma in extra:
    lech.append(ma)
    wk.append([ma, plan_cnt[ma], None, None, None, None, None, None, "LỆCH: mã có trên mặt bằng, không có trong bảng"])
n = wk.max_row
wk.append(["Tổng", f"=SUM(B2:B{n})", f"=SUM(C2:C{n})", f"=SUM(D2:D{n})"])

# Sheet gia thiet
wg = wb.create_sheet("Gia thiet")
wg.append(["STT", "Giả định / ghi chú"])
for c in wg[1]:
    c.font = B
for i, g in enumerate(GT, 1):
    wg.append([i, g])
wg.column_dimensions["B"].width = 140
wb.save(OUT)

# In tom tat
def tot(g, dv):
    return sum(r["kl"] for r in rows if r["nhom"] == g and r["dv"] == dv and ("sub" in r or "vol_of" in r or g != 1))
print("Tuong 200 m2:", tot(1, "m2"), " m3:", tot(1, "m3"))
for k, *_ in TANG:
    print("  ", k, [r["kl"] for r in rows if r.get("ma") == f"TX-{k}" and "sub" in r])
print("BT san m3:", tot(2, "m3"))
for k, *_ in SLABS:
    print("  ", k, sum(r["kl"] for r in rows if r["ma"] == f"BTS-{k}"))
print("BT cot m3:", tot(3, "m3"))
print("Cua bo:", tot(4, "bộ"), " m2:", tot(4, "m2"))
print("Lech cua:", lech or "khong")
print("Saved", OUT)
